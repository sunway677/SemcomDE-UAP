"""
Evaluation of the data-efficient SemCom UAPs: CR-ASR at Recall@1 under AWGN.

Implements the evaluation protocol of the paper (Sec. II system model,
Sec. IV-A "Dataset and Metric") on the Flickr30K 1K test split
(configs/Retrieval_flickr_test.yaml: 1,000 images, 5,000 captions), with the
source (surrogate) and target VLP models drawn from ALBEF, TCL and BLIP.

Perturbations. The frozen UAPs written by train_data_efficient_uap.py are
read from <load_dir>/<SOURCE>/<dataset>/:
    uap_noise-<epoch>.pth      image UAP delta_x, float32 (1, 3, H, W) with
                               values in [-eps_x, eps_x], H = W = image_res
                               ((3, H, W) is accepted as well);
    uap_embedding-<epoch>.pth  text UAP delta_t, float32 (eps_t, D), already
                               norm-clamped (eps_t = --text_budget).
They are "the final generated image and text UAPs" of Sec. III-E: the
delta_x / delta_t produced in the last optimization step of epoch <epoch>,
at the last sampled SNR stratum of that step. The default --epoch 49 is the
last epoch of the 50-epoch schedule of Sec. IV-A.

Adversarial queries (Sec. II "Adversary's Goal", Sec. III-B.4, Sec. III-C).
    Image: x + delta_x, with the full 3-channel delta_x broadcast over the
        batch and the result kept in the valid pixel range [0, 1], as in
        training.
    Text: T (+) delta_t. Each row of delta_t is mapped to its nearest
        vocabulary embedding (L2 distance) among the candidate tokens of
        word_idx_<SOURCE>.npy in word_embedding_<SOURCE>.pth, skipping filter
        words (the BERT-based models fall back to the ALBEF files when the
        per-model files are missing). TextAttacker then substitutes the
        resulting word(s) for the most important word(s) of every caption,
        with word importance measured on the source model. The adversarial
        captions are built ONCE with the source model and the same captions
        are used for every target model, so the result does not depend on
        the order of --target_models.

Galleries. The I2T text gallery holds the clean captions and the T2I image
gallery holds the clean images: no UAP and no channel noise (the adversary
perturbs only the queries and has no access to the retrieval database,
Sec. II).

Channel (Sec. II). A query embedding e (L2-normalized, dimension d) is
received as
    e~ = norm(e + n),   n ~ N(0, sigma^2(gamma) I),
    sigma^2(gamma) = 1 / (d * 10^(gamma / 10)),
for an SNR gamma in dB; gamma = inf denotes a noiseless channel (e~ = e).
Retrieval ranks the gallery of the opposite modality by cosine similarity.

Metric (Sec. IV-A). Clean-Retrieval Attack Success Rate at Recall@1:
    CR-ASR(gamma) = |S(gamma) & F(inf)| / |S(gamma)|,
    S(gamma): clean queries retrieved correctly at R@1 when the clean query
              embedding is transmitted over the AWGN channel at SNR gamma;
    F(inf):   adversarial queries (UAP applied) that fail at R@1 without
              channel noise.
A query succeeds at R@1 if, for I2T, a ground-truth caption of the query
image is ranked first, and, for T2I, the ground-truth image of the query
caption is ranked first.

For every target model, retrieval direction and SNR in --snr_list the script
reports the clean Recall@1 at gamma, the adversarial Recall@1 at gamma = inf
and CR-ASR(gamma), all in percent, prints them as a table and writes the same
numbers to outputs/<output>/<SOURCE>_<dataset>_cr_asr.xlsx.

Usage:
python eval_data_efficient_uap.py \
    --config ./configs/Retrieval_flickr_test.yaml \
    --source_model ALBEF \
    --checkpoint ./checkpoint \
    --load_dir ./output_n500_ALBEF \
    --epoch 49 \
    --snr_list inf 20 15 10 5 0 \
    --output result_n500_ALBEF
"""

import argparse
import gc
import math
import os
import random
import zlib

import numpy as np
import torch
import torch.nn.functional as F
import torch.backends.cudnn as cudnn
from openpyxl import Workbook
from PIL import Image
from ruamel.yaml import YAML
from torch.utils.data import DataLoader
from torchvision import transforms
from tqdm import tqdm

from utils import load_model, get_filter_words
from attacker import TextAttacker
from dataset import paired_dataset

EVAL_MODELS  = ['ALBEF', 'TCL', 'BLIP']
MAX_TEXT_LEN = 30
DIRECTIONS   = ('I2T', 'T2I')

# Same input normalization as in training.
IMAGE_NORMALIZE = transforms.Normalize(
    (0.48145466, 0.4578275, 0.40821073),
    (0.26862954, 0.26130258, 0.27577711))


# ==========================================
# Arguments
# ==========================================

def parse_snr(value):
    """Parse one --snr_list entry: a finite SNR in dB, or 'inf' (noiseless)."""
    if value.strip().lower() in ('inf', '+inf'):
        return float('inf')
    try:
        snr = float(value)
    except ValueError:
        raise argparse.ArgumentTypeError(f'invalid SNR {value!r}: expected a number (dB) or "inf"')
    if not math.isfinite(snr):
        raise argparse.ArgumentTypeError(f'invalid SNR {value!r}')
    return snr


def snr_tag(snr_db):
    return 'inf' if math.isinf(snr_db) else f'{snr_db:g}'


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description='CR-ASR evaluation of data-efficient SemCom UAPs on '
                    'I2T / T2I retrieval under an AWGN channel (paper Sec. II, IV-A).')
    parser.add_argument('--config', default='./configs/Retrieval_flickr_test.yaml',
                        help='test-split config (annotation file, image root, image_res)')
    parser.add_argument('--checkpoint', default='./checkpoint',
                        help='directory holding <MODEL>/<dataset>.pth fine-tuned retrieval checkpoints')
    parser.add_argument('--text_encoder', default='bert-base-uncased',
                        help='HuggingFace BERT name used for the tokenizer and TextAttacker reference model')
    parser.add_argument('--source_model', default='ALBEF', choices=EVAL_MODELS,
                        help='surrogate model the UAPs were trained on; also builds the adversarial captions')
    parser.add_argument('--target_models', nargs='+', default=list(EVAL_MODELS), choices=EVAL_MODELS,
                        help='victim models to evaluate')
    parser.add_argument('--load_dir', required=True,
                        help='training --save_dir; UAPs are read from <load_dir>/<SOURCE>/<dataset>/')
    parser.add_argument('--epoch', type=int, default=49,
                        help='epoch index of the UAP files (default: 49, the last of 50 epochs)')
    parser.add_argument('--text_budget', '--word_num', dest='text_budget', type=int, default=1,
                        help='text budget eps_t in tokens; must equal the number of rows of the text UAP')
    parser.add_argument('--snr_list', nargs='+', type=parse_snr,
                        default=[float('inf'), 20.0, 15.0, 10.0, 5.0, 0.0],
                        help='channel SNRs gamma in dB; "inf" denotes a noiseless channel')
    parser.add_argument('--batch_size', type=int, default=4,
                        help='images per encoding batch; captions are encoded in batches of 5x this')
    parser.add_argument('--seed', type=int, default=42,
                        help='seed for all random number generators, including the channel noise')
    parser.add_argument('--output', default='result_data_efficient_uap',
                        help='results are written to outputs/<output>/')
    args = parser.parse_args(argv)
    # Keep the given order but drop duplicates.
    args.target_models = list(dict.fromkeys(args.target_models))
    args.snr_list = list(dict.fromkeys(args.snr_list))
    if args.text_budget < 1:
        parser.error('--text_budget must be >= 1')
    return args


def seed_everything(seed):
    torch.manual_seed(seed)
    np.random.seed(seed)
    random.seed(seed)
    cudnn.benchmark = False
    cudnn.deterministic = True


# ==========================================
# UAP loading
# ==========================================

def load_image_uap(path, image_res):
    """Load the image UAP delta_x as a float32 (1, 3, H, W) tensor."""
    if not os.path.exists(path):
        raise FileNotFoundError(f'image UAP not found: {path}')
    delta_x = torch.load(path, map_location='cpu', weights_only=True).float()
    if delta_x.dim() == 3:
        delta_x = delta_x.unsqueeze(0)
    if delta_x.dim() != 4 or delta_x.shape[0] != 1 or delta_x.shape[1] != 3:
        raise ValueError(f'{path}: expected an image UAP of shape (1, 3, H, W) or (3, H, W), '
                         f'got {tuple(delta_x.shape)}')
    if tuple(delta_x.shape[-2:]) != (image_res, image_res):
        raise ValueError(f'{path}: spatial size {tuple(delta_x.shape[-2:])} does not match '
                         f'image_res={image_res} of the config')
    return delta_x


def load_text_uap(path, text_budget):
    """Load the text UAP delta_t as a float32 (text_budget, D) tensor."""
    if not os.path.exists(path):
        raise FileNotFoundError(f'text UAP not found: {path}')
    delta_t = torch.load(path, map_location='cpu', weights_only=True).float()
    if delta_t.dim() == 1:
        delta_t = delta_t.unsqueeze(0)
    if delta_t.dim() != 2 or delta_t.shape[0] != text_budget:
        raise ValueError(f'{path}: expected a text UAP of shape ({text_budget}, D) '
                         f'(--text_budget {text_budget}), got {tuple(delta_t.shape)}')
    return delta_t


# ==========================================
# Text UAP -> adversarial captions
# ==========================================

def resolve_vocab_file(pattern, source_model):
    """Return the per-model vocabulary file, falling back to the ALBEF one.

    All BERT-based models share the ALBEF word-embedding table and candidate
    vocabulary (scripts/generate_word_embeddings.py writes the ALBEF files and
    aliases them for the other models).
    """
    names = [source_model] + (['ALBEF'] if source_model != 'ALBEF' else [])
    for name in names:
        path = pattern.format(name)
        if os.path.exists(path):
            return path
    raise FileNotFoundError(
        f'none of {[pattern.format(n) for n in names]} found in {os.getcwd()}; '
        f'run from the repo root after "python scripts/generate_word_embeddings.py --all"')


def text_uap_to_adv_words(delta_t, word_embeddings, word_idxs, tokenizer, filter_words):
    """Map each row of delta_t to its nearest admissible vocabulary word.

    Realizes the mapping of delta_t to valid token substitutions
    (Sec. III-B.4b): the candidates are the tokens listed in word_idxs, the
    distance is L2 in the word-embedding space, and filter words are skipped.
    """
    word_idxs  = np.asarray(word_idxs, dtype=np.int64)
    candidates = word_embeddings[torch.from_numpy(word_idxs)].float()
    adv_words  = []
    for row in delta_t.float():
        dist = torch.linalg.vector_norm(candidates - row.unsqueeze(0), dim=1)
        for k in torch.argsort(dist, stable=True).tolist():
            word = tokenizer.convert_ids_to_tokens(int(word_idxs[k]))
            if word not in filter_words:
                adv_words.append(word)
                break
        else:
            raise ValueError('no admissible vocabulary word for the text UAP')
    return adv_words


@torch.no_grad()
def build_adv_captions(txt_attacker, source_model, texts, chunk_size):
    """Apply the text UAP to every caption, T_adv = T (+) delta_t.

    Word importance is measured on the source model. The returned list is
    aligned with `texts` and is reused unchanged for every target model.
    """
    adv_texts = []
    for start in tqdm(range(0, len(texts), chunk_size), desc='Adversarial captions'):
        adv_texts.extend(txt_attacker.get_adv_text(source_model, texts[start:start + chunk_size]))
    if len(adv_texts) != len(texts):
        raise RuntimeError('TextAttacker returned a different number of captions')
    return adv_texts


# ==========================================
# Encoders (semantic embeddings, Eq. 1 / Eq. 6)
# ==========================================

@torch.no_grad()
def encode_images(model, data_loader, device, delta_x=None):
    """L2-normalized image embeddings in dataset order.

    delta_x=None gives the clean embeddings (clean queries and the T2I
    gallery); otherwise the UAP is added to every image first (adversarial
    queries x + delta_x).
    """
    feats = None
    for images, _, img_ids, _ in tqdm(data_loader, desc='  images' + ('' if delta_x is None else ' + UAP')):
        images = images.to(device)
        if delta_x is not None:
            images = torch.clamp(images + delta_x.to(device), 0.0, 1.0)
        out = model.inference_image(IMAGE_NORMALIZE(images))['image_feat'].float().cpu()
        if feats is None:
            feats = torch.zeros(len(data_loader.dataset), out.shape[-1])
        feats[img_ids] = out
    return F.normalize(feats, dim=-1)


@torch.no_grad()
def encode_texts(model, tokenizer, texts, device, batch_size):
    """L2-normalized text embeddings, aligned with `texts`."""
    feats = []
    for start in tqdm(range(0, len(texts), batch_size), desc='  texts'):
        text_input = tokenizer(texts[start:start + batch_size], padding='max_length',
                               truncation=True, max_length=MAX_TEXT_LEN,
                               return_tensors='pt').to(device)
        feats.append(model.inference_text(text_input)['text_feat'].float().cpu())
    return F.normalize(torch.cat(feats, dim=0), dim=-1)


# ==========================================
# AWGN channel (Sec. II)
# ==========================================

def noise_std(dim, snr_db):
    """sigma(gamma) = sqrt(1 / (d * 10^(gamma/10))) (Sec. II)."""
    return math.sqrt(1.0 / (dim * 10.0 ** (snr_db / 10.0)))


def channel_generator(seed, model_name, modality, snr_db):
    """Dedicated noise generator per (target model, query modality, SNR).

    The seed is derived from a stable hash, so a noise realization does not
    depend on the order of --target_models or --snr_list.
    """
    key = f'{seed}|{model_name}|{modality}|{snr_tag(snr_db)}'
    return torch.Generator().manual_seed(zlib.crc32(key.encode('utf-8')))


def awgn_channel(feats, snr_db, generator):
    """Transmit embeddings: e~ = norm(e + n), n ~ N(0, sigma^2(gamma) I)."""
    if math.isinf(snr_db):
        return feats
    noise = torch.randn(feats.shape, generator=generator, dtype=feats.dtype)
    return F.normalize(feats + noise * noise_std(feats.shape[-1], snr_db), dim=-1)


# ==========================================
# Retrieval and metrics
# ==========================================

def r1_success_i2t(img_queries, txt_gallery, txt2img):
    """I2T success at R@1: a ground-truth caption of the query image is ranked first.

    Every caption belongs to exactly one image, so the top-1 caption is a
    ground-truth caption iff txt2img[top1] is the query image.
    """
    top1 = (img_queries @ txt_gallery.t()).argmax(dim=1)
    return txt2img[top1] == torch.arange(img_queries.shape[0])


def r1_success_t2i(txt_queries, img_gallery, txt2img):
    """T2I success at R@1: the ground-truth image of the query caption is ranked first."""
    top1 = (txt_queries @ img_gallery.t()).argmax(dim=1)
    return top1 == txt2img


def cr_asr(clean_success, adv_success_inf):
    """CR-ASR = |S & F| / |S| in percent (NaN when S is empty).

    clean_success:   boolean mask of S(gamma);
    adv_success_inf: boolean mask of adversarial R@1 success at gamma = inf,
                     whose complement is F(inf).
    """
    n_s = int(clean_success.sum())
    if n_s == 0:
        return float('nan')
    return 100.0 * int((clean_success & ~adv_success_inf).sum()) / n_s


def compute_metrics(img_clean, txt_clean, img_adv, txt_adv, txt2img,
                    snr_list, seed, model_name):
    """Clean R@1, adversarial R@1 and CR-ASR for one target model.

    img_clean / txt_clean: clean embeddings; they form the clean galleries
        (never perturbed or noised) and, after the channel, the clean queries.
    img_adv / txt_adv: adversarial query embeddings, retrieved without
        channel noise to obtain F(inf).
    Returns {direction: {snr: {'clean_r1', 'adv_r1_inf', 'cr_asr'}}}.
    """
    adv_ok = {
        'I2T': r1_success_i2t(img_adv, txt_clean, txt2img),
        'T2I': r1_success_t2i(txt_adv, img_clean, txt2img),
    }
    metrics = {d: {} for d in DIRECTIONS}
    for snr_db in snr_list:
        img_rx = awgn_channel(img_clean, snr_db, channel_generator(seed, model_name, 'image', snr_db))
        txt_rx = awgn_channel(txt_clean, snr_db, channel_generator(seed, model_name, 'text', snr_db))
        clean_ok = {
            'I2T': r1_success_i2t(img_rx, txt_clean, txt2img),
            'T2I': r1_success_t2i(txt_rx, img_clean, txt2img),
        }
        for d in DIRECTIONS:
            metrics[d][snr_db] = {
                'clean_r1':   100.0 * clean_ok[d].float().mean().item(),
                'adv_r1_inf': 100.0 * adv_ok[d].float().mean().item(),
                'cr_asr':     cr_asr(clean_ok[d], adv_ok[d]),
            }
    return metrics


# ==========================================
# Models
# ==========================================

def checkpoint_path(args, model_name, config):
    return os.path.join(args.checkpoint, model_name, f'{config["dataset"]}.pth')


def load_eval_model(args, model_name, config, device):
    """Load a retrieval model (eval mode, on `device`), its reference MLM and tokenizer."""
    print(f'Loading {model_name}...')
    model, ref_model, tokenizer = load_model(
        model_name, checkpoint_path(args, model_name, config), args.text_encoder,
        config, torch.device('cpu'))
    return model.to(device).eval(), ref_model, tokenizer


def free_memory():
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()


# ==========================================
# Reporting
# ==========================================

def fmt_pct(value):
    return 'nan' if math.isnan(value) else f'{value:.2f}'


def result_rows(results, snr_list):
    """Flatten results into (target, direction, snr, clean, adv, cr_asr) rows, rounded to 0.01."""
    rows = []
    for target, metrics in results.items():
        for d in DIRECTIONS:
            for snr_db in snr_list:
                m = metrics[d][snr_db]
                rows.append((target, d, snr_tag(snr_db),
                             round(m['clean_r1'], 2), round(m['adv_r1_inf'], 2),
                             round(m['cr_asr'], 2)))
    return rows


def print_results(source_model, adv_words, rows, snr_list):
    print(f'\n{"=" * 72}')
    print(f'Source {source_model} | adversarial word(s) {adv_words} | values in %')
    print('CR-ASR(gamma) = |S(gamma) & F(inf)| / |S(gamma)| at Recall@1')
    print(f'{"=" * 72}')
    print(f'{"Target":<8} {"Dir":<5} {"SNR(dB)":>8} {"Clean R@1":>10} {"Adv R@1 (inf)":>14} {"CR-ASR":>8}')
    print('-' * 72)
    for target, d, tag, clean, adv, asr in rows:
        print(f'{target:<8} {d:<5} {tag:>8} {fmt_pct(clean):>10} {fmt_pct(adv):>14} {fmt_pct(asr):>8}')

    # Same CR-ASR numbers, one line per target and direction (Table I layout).
    tags = [snr_tag(s) for s in snr_list]
    print(f'\nCR-ASR (%) by SNR, source {source_model}')
    print(f'{"Target":<8} {"Dir":<5} ' + ' '.join(f'{t:>8}' for t in tags))
    by_key = {}
    for target, d, tag, _, _, asr in rows:
        by_key.setdefault((target, d), {})[tag] = asr
    for (target, d), values in by_key.items():
        print(f'{target:<8} {d:<5} ' + ' '.join(f'{fmt_pct(values[t]):>8}' for t in tags))


def xlsx_value(value):
    return None if isinstance(value, float) and math.isnan(value) else value


def write_xlsx(path, source_model, rows, snr_list, settings):
    wb = Workbook()
    ws = wb.active
    ws.title = 'CR-ASR'
    ws.append(['Source', 'Target', 'Direction', 'SNR (dB)', 'Clean R@1', 'Adv R@1 (inf)', 'CR-ASR'])
    for target, d, tag, clean, adv, asr in rows:
        ws.append([source_model, target, d, tag, xlsx_value(clean), xlsx_value(adv), xlsx_value(asr)])

    tags = [snr_tag(s) for s in snr_list]
    ws_t = wb.create_sheet('CR-ASR by SNR')
    ws_t.append(['Source', 'Target', 'Direction'] + [f'CR-ASR {t} dB' if t != 'inf' else 'CR-ASR inf'
                                                     for t in tags])
    by_key = {}
    for target, d, tag, _, _, asr in rows:
        by_key.setdefault((target, d), {})[tag] = asr
    for (target, d), values in by_key.items():
        ws_t.append([source_model, target, d] + [xlsx_value(values[t]) for t in tags])

    ws_s = wb.create_sheet('Settings')
    for key, value in settings:
        ws_s.append([key, value])
    wb.save(path)


# ==========================================
# Entry point
# ==========================================

def main(argv=None):
    args = parse_args(argv)
    seed_everything(args.seed)
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')

    with open(args.config, 'r') as f:
        config = YAML(typ='safe').load(f)

    for name in dict.fromkeys([args.source_model] + args.target_models):
        ckpt = checkpoint_path(args, name, config)
        if not os.path.exists(ckpt):
            raise FileNotFoundError(f'{name} checkpoint not found: {ckpt}')

    # Test split; same deterministic preprocessing as training.
    test_transform = transforms.Compose([
        transforms.Resize((config['image_res'], config['image_res']),
                          interpolation=Image.BICUBIC),
        transforms.ToTensor()])
    test_dataset = paired_dataset(config['annotation_file'], test_transform, config['image_root'])
    test_loader  = DataLoader(test_dataset, batch_size=args.batch_size, shuffle=False,
                              num_workers=0, collate_fn=test_dataset.collate_fn)
    clean_texts = list(test_dataset.text)
    txt2img     = torch.tensor([test_dataset.txt2img[j] for j in range(len(clean_texts))],
                               dtype=torch.long)
    text_batch  = 5 * args.batch_size
    print(f'Test set: {len(test_dataset)} images, {len(clean_texts)} captions')

    # Frozen UAPs: the final generated image and text UAPs of training (Sec. III-E).
    load_dir   = os.path.join(args.load_dir, args.source_model, config['dataset'])
    noise_path = os.path.join(load_dir, f'uap_noise-{args.epoch}.pth')
    emb_path   = os.path.join(load_dir, f'uap_embedding-{args.epoch}.pth')
    delta_x = load_image_uap(noise_path, config['image_res'])
    delta_t = load_text_uap(emb_path, args.text_budget)
    print(f'Image UAP {tuple(delta_x.shape)}, ||delta_x||_inf = {delta_x.abs().max().item():.6f} '
          f'({delta_x.abs().max().item() * 255:.2f}/255)')
    print(f'Text UAP  {tuple(delta_t.shape)}')

    # Adversarial captions: built once with the source model, reused for all targets.
    we_path  = resolve_vocab_file('word_embedding_{}.pth', args.source_model)
    idx_path = resolve_vocab_file('word_idx_{}.npy', args.source_model)
    src_model, ref_model, src_tokenizer = load_eval_model(args, args.source_model, config, device)
    adv_words = text_uap_to_adv_words(
        delta_t, torch.load(we_path, map_location='cpu', weights_only=True),
        np.load(idx_path), src_tokenizer, get_filter_words())
    print(f'Adversarial word(s): {adv_words} (vocabulary: {we_path}, {idx_path})')
    txt_attacker = TextAttacker(
        ref_model.to(device), src_tokenizer, cls=False, max_length=MAX_TEXT_LEN,
        number_perturbation=args.text_budget, adv_words=adv_words,
        topk=10, threshold_pred_score=0.3)
    adv_texts = build_adv_captions(txt_attacker, src_model, clean_texts, text_batch)
    print(f'Example: "{clean_texts[0]}" -> "{adv_texts[0]}"')
    del src_model, ref_model, txt_attacker
    free_memory()

    snr_list = args.snr_list
    print(f'SNR list (dB): {[snr_tag(s) for s in snr_list]}')

    results = {}
    for target in args.target_models:
        model, _, tokenizer = load_eval_model(args, target, config, device)
        print(f'[{target}] encoding clean / adversarial queries and clean galleries')
        img_clean = encode_images(model, test_loader, device)
        img_adv   = encode_images(model, test_loader, device, delta_x=delta_x)
        txt_clean = encode_texts(model, tokenizer, clean_texts, device, text_batch)
        txt_adv   = encode_texts(model, tokenizer, adv_texts, device, text_batch)
        results[target] = compute_metrics(img_clean, txt_clean, img_adv, txt_adv,
                                          txt2img, snr_list, args.seed, target)
        del model, tokenizer, img_clean, img_adv, txt_clean, txt_adv
        free_memory()

    rows = result_rows(results, snr_list)
    print_results(args.source_model, adv_words, rows, snr_list)

    save_dir = os.path.join('outputs', args.output)
    os.makedirs(save_dir, exist_ok=True)
    xlsx_path = os.path.join(save_dir, f'{args.source_model}_{config["dataset"]}_cr_asr.xlsx')
    settings = [
        ('source_model', args.source_model),
        ('target_models', ' '.join(args.target_models)),
        ('image_uap', noise_path),
        ('text_uap', emb_path),
        ('adversarial_words', ' '.join(adv_words)),
        ('snr_list_db', ' '.join(snr_tag(s) for s in snr_list)),
        ('seed', args.seed),
        ('values', 'percent'),
        ('CR-ASR', 'CR-ASR(gamma) = |S(gamma) & F(inf)| / |S(gamma)| at Recall@1'),
        ('S(gamma)', 'clean queries correct at R@1 when the clean query embedding is sent over AWGN at SNR gamma'),
        ('F(inf)', 'adversarial queries (UAP applied) that fail at R@1 without channel noise'),
        ('galleries', 'clean images (T2I) and clean captions (I2T): no UAP, no channel noise'),
    ]
    write_xlsx(xlsx_path, args.source_model, rows, snr_list, settings)
    print(f'\nSaved to {xlsx_path}')


if __name__ == '__main__':
    main()
